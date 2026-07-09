import fitz
from docx import Document
import os


def extract_pdf_text(file):

    text = ""

    # If a file path (string) is given
    if isinstance(file, str):
        pdf = fitz.open(file)

    # If a Streamlit uploaded file is given
    else:
        pdf = fitz.open(stream=file.read(), filetype="pdf")

    for page in pdf:
        text += page.get_text()

    return text


def extract_docx_text(file):

    text = ""

    # If a file path (string)
    if isinstance(file, str):
        document = Document(file)

    # If uploaded file
    else:
        document = Document(file)

    for para in document.paragraphs:
        text += para.text + "\n"

    return text


def extract_resume_text(file):

    # File path
    if isinstance(file, str):
        extension = os.path.splitext(file)[1].lower()

    # Uploaded file
    else:
        extension = os.path.splitext(file.name)[1].lower()

    if extension == ".pdf":
        return extract_pdf_text(file)

    elif extension == ".docx":
        return extract_docx_text(file)

    return ""